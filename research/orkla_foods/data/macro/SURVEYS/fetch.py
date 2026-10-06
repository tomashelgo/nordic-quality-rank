#!/usr/bin/env python3
"""
fetch.py -- Business & consumer SURVEY series for the Orkla Foods predictor analysis.

Writes one CSV per series (columns date,value; date = first day of the period) plus
catalog.csv into the directory this script lives in.

Usage:
    python3 -I fetch.py [--download-dir DIR]

All raw downloads go to DIR (default: a fresh temporary directory). Nothing is executed
from the download directory; downloaded files are only parsed as data.

Sources (all free, no API keys, no manual steps required as of 2026-10-06):
  1. European Commission DG ECFIN Business and Consumer Surveys (BCS) -- zip files of
     seasonally adjusted time series (main indicators, consumer, retail, industry and
     NACE-2 industry/retail sub-sector files). The current folder name
     (".../series/nace2_ecfin_YYMM/") is scraped from the EC "time series" page.
  2. OECD SDMX REST API, dataflow OECD.SDD.STES,DSD_STES@DF_CLI (CLI, CCI, BCI;
     amplitude adjusted).  Note: since the 2023 CLI redesign the OECD publishes the CLI
     only for G7/G20-type economies; NOR/SWE/DNK/FIN/CZE/AUT/POL CLIs no longer exist.
  3. University of Michigan Surveys of Consumers (sca.isr.umich.edu, table tbmics.csv),
     with FRED UMCSENT as fallback.
  4. Konjunkturinstitutet / NIER PxWeb API (statistik.konj.se) -- Economic Tendency Survey.
  5. Statistics Norway (SSB) PxWeb API -- business tendency survey tables 08267 and 08264.
  6. Norges Bank Regional Network: latest report xlsx (scraped from the reports list page).
  7. Norges Bank Expectations Survey: latest xlsx (scraped from the topic page).
  8. Finans Norge Forventningsbarometer (with Verian/Kantar): data are only published as
     embedded Datawrapper charts on finansnorge.no; the script scrapes the chart ids from
     the page, resolves the latest published chart version and downloads dataset.csv.
     (Chart ids at time of writing: dPMwa = main indicator, 4EDKP = sub-indicators.)

Manual notes:
  * If a source page layout changes, the scrapers raise a clear error for that source;
    other sources still run (each source is wrapped in try/except and failures are
    printed at the end).
  * Norway PMI (DNB/NIMA) is not freely available as data and was discontinued after the
    Nov-2025 release -- not collected.
"""
from __future__ import annotations

import argparse
import csv
import io
import itertools
import json
import re
import sys
import tempfile
import time
import traceback
import zipfile
from datetime import date
from pathlib import Path

import numpy as np
import pandas as pd
import requests

OUT = Path(__file__).resolve().parent
UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                    "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"}

CATALOG_COLS = ["series_id", "file", "description", "country", "frequency", "unit",
                "seasonal_adjustment", "start", "end", "n_obs", "source", "source_query",
                "approx_release_lag_days", "notes"]

RESULTS: list[dict] = []      # catalog rows
SERIES: dict[str, pd.Series] = {}
FAILURES: list[tuple[str, str]] = []


# ----------------------------------------------------------------------------- helpers
def http_get(url, dl_dir: Path | None = None, fname: str | None = None, headers=None,
             tries=3, timeout=120):
    last = None
    for i in range(tries):
        try:
            r = requests.get(url, headers=headers if headers is not None else UA, timeout=timeout)
            if r.status_code == 200:
                if dl_dir is not None and fname:
                    (dl_dir / fname).write_bytes(r.content)
                return r
            last = f"HTTP {r.status_code}"
            if r.status_code in (403, 404, 407):
                break
        except Exception as e:  # noqa: BLE001
            last = repr(e)
        time.sleep(2 * (i + 1))
    raise RuntimeError(f"GET failed {url}: {last}")


def to_month_start(idx) -> pd.DatetimeIndex:
    return pd.DatetimeIndex(pd.to_datetime(idx)).to_period("M").to_timestamp()


def quarter_label_to_date(lbl: str) -> pd.Timestamp | None:
    """Parse '1992-Q3', '2026K2', '3. q. 2026', '2.q. 2023' -> first day of quarter."""
    s = str(lbl).strip()
    m = re.match(r"^(\d{4})\s*[-]?\s*[QK](\d)$", s, re.I)
    if m:
        y, q = int(m.group(1)), int(m.group(2))
    else:
        m = re.match(r"^(\d)\s*\.\s*q\s*\.?\s*(\d{4})$", s, re.I)
        if not m:
            return None
        q, y = int(m.group(1)), int(m.group(2))
    return pd.Timestamp(y, 3 * (q - 1) + 1, 1)


def add(series_id, s: pd.Series, *, description, country, frequency, unit, sa, source,
        source_query, lag, notes=""):
    s = pd.to_numeric(s, errors="coerce").dropna()
    s = s[~s.index.duplicated(keep="last")].sort_index()
    if s.empty:
        FAILURES.append((series_id, "empty after parsing"))
        return
    s.index = pd.DatetimeIndex(s.index)
    # document internal gaps automatically
    step = {"M": 1, "Q": 3}.get(frequency)
    if step:
        full = pd.date_range(s.index.min(), s.index.max(), freq="MS")[::step]
        miss = full.difference(s.index)
        if len(miss):
            # compress to ranges
            rngs, start, prev = [], miss[0], miss[0]
            for d in list(miss[1:]) + [None]:
                if d is not None and (d.year * 12 + d.month) - (prev.year * 12 + prev.month) == step:
                    prev = d
                    continue
                rngs.append(start.strftime("%Y-%m") + ("" if start == prev else ".." + prev.strftime("%Y-%m")))
                if d is not None:
                    start = prev = d
            notes = (notes + " " if notes else "") + f"Internal gap(s) in source data: {', '.join(rngs)}."
    SERIES[series_id] = s
    RESULTS.append(dict(series_id=series_id, file=f"{series_id}.csv", description=description,
                        country=country, frequency=frequency, unit=unit,
                        seasonal_adjustment=sa, start=s.index.min().strftime("%Y-%m-%d"),
                        end=s.index.max().strftime("%Y-%m-%d"), n_obs=int(s.size), source=source,
                        source_query=source_query, approx_release_lag_days=lag, notes=notes))


def write_all():
    for sid, s in SERIES.items():
        df = pd.DataFrame({"date": s.index.strftime("%Y-%m-%d"), "value": np.round(s.values.astype(float), 6)})
        df.to_csv(OUT / f"{sid}.csv", index=False)
    cat = pd.DataFrame(RESULTS, columns=CATALOG_COLS).sort_values("series_id")
    cat.to_csv(OUT / "catalog.csv", index=False, quoting=csv.QUOTE_MINIMAL)


def run_source(name, fn, *args):
    print(f"--- {name}", flush=True)
    try:
        fn(*args)
    except Exception as e:  # noqa: BLE001
        traceback.print_exc()
        FAILURES.append((name, repr(e)))


# ----------------------------------------------------------------------------- 1. EC BCS
EC_PAGE = ("https://economy-finance.ec.europa.eu/economic-forecast-and-surveys/"
           "business-and-consumer-surveys/download-business-and-consumer-survey-data/time-series_en")
EC_SUB_PAGE = ("https://economy-finance.ec.europa.eu/economic-forecast-and-surveys/"
               "business-and-consumer-surveys/download-business-and-consumer-survey-data/subsector-data_en")
EC_BASE = "https://ec.europa.eu/economy_finance/db_indicators/surveys/documents/series/"
EC_COUNTRIES = {"DK": "Denmark", "FI": "Finland", "EE": "Estonia", "LV": "Latvia",
                "LT": "Lithuania", "CZ": "Czechia", "SK": "Slovakia", "AT": "Austria",
                "PL": "Poland", "HU": "Hungary", "RO": "Romania", "DE": "Germany",
                "SE": "Sweden", "EA": "Euro area", "EU": "European Union"}


EC_ZIP_OF = {"main_indicators_nace2.xlsx": "main_indicators_sa_nace2",
             "consumer_total_sa_nace2.xlsx": "consumer_total_sa_nace2",
             "retail_total_sa_nace2.xlsx": "retail_total_sa_nace2",
             "industry_total_sa_nace2.xlsx": "industry_total_sa_nace2",
             "industry_subsectors_sa_m_nace2.xlsx": "industry_subsectors_sa_nace2",
             "retail_subsectors_sa_m_nace2.xlsx": "retail_subsectors_sa_nace2"}


def ec_folder() -> str:
    folders = set()
    for url in (EC_PAGE, EC_SUB_PAGE):
        try:
            html = http_get(url).text
            folders |= set(re.findall(r"series/(nace2_ecfin_\d{4})/", html))
        except Exception as e:  # noqa: BLE001
            print("EC page scrape failed:", e)
    if folders:
        return max(folders)
    # fallback: probe the last 4 months
    t = date.today()
    for k in range(4):
        y, m = t.year, t.month - k
        while m <= 0:
            y, m = y - 1, m + 12
        f = f"nace2_ecfin_{y % 100:02d}{m:02d}"
        r = requests.head(EC_BASE + f + "/main_indicators_sa_nace2.zip", headers=UA, timeout=60)
        if r.status_code == 200:
            return f
    raise RuntimeError("could not determine EC BCS folder")


def ec_sheet(xdir: Path, fname: str, sheet: str) -> pd.DataFrame:
    df = pd.read_excel(xdir / fname, sheet_name=sheet)
    df = df.set_index(df.columns[0])
    df.index = pd.to_datetime(df.index, errors="coerce")
    df = df[df.index.notna()]
    df.index = to_month_start(df.index)
    return df


def fetch_ec(dl: Path):
    folder = ec_folder()
    print("EC folder:", folder)
    zips = ["main_indicators_sa_nace2", "consumer_total_sa_nace2", "retail_total_sa_nace2",
            "industry_total_sa_nace2", "industry_subsectors_sa_nace2", "retail_subsectors_sa_nace2"]
    xdir = dl / "ec_extracted"
    xdir.mkdir(exist_ok=True)
    for z in zips:
        http_get(f"{EC_BASE}{folder}/{z}.zip", dl, f"{z}.zip", timeout=300)
        with zipfile.ZipFile(dl / f"{z}.zip") as zf:
            for m in zf.infolist():
                if m.filename.endswith(".xlsx") and "/" not in m.filename:
                    zf.extract(m, xdir)
    src = "European Commission DG ECFIN Business and Consumer Surveys (BCS), seasonally adjusted"
    lag_note = ("Released on the penultimate working day of the reference month "
                "(flash consumer confidence for EU/EA ~9 days before month end)")
    specs = [
        # (file, sheet, column pattern, id suffix, description, unit)
        ("main_indicators_nace2.xlsx", "MONTHLY", "{c}.ESI", "ec_esi",
         "Economic Sentiment Indicator (ESI)", "index, long-term average=100"),
        ("consumer_total_sa_nace2.xlsx", "CONSUMER MONTHLY", "CONS.{c}.TOT.COF.BS.M", "ec_cons_conf",
         "Consumer confidence indicator (avg of Q1,Q2,Q4,Q9)", "balance, % points"),
        ("consumer_total_sa_nace2.xlsx", "CONSUMER MONTHLY", "CONS.{c}.TOT.6.BS.M", "ec_cons_price_exp_12m",
         "Consumers: price trends over next 12 months (Q6)", "balance, % points"),
        ("consumer_total_sa_nace2.xlsx", "CONSUMER MONTHLY", "CONS.{c}.TOT.5.BS.M", "ec_cons_price_past_12m",
         "Consumers: perceived price trends over last 12 months (Q5)", "balance, % points"),
        ("retail_total_sa_nace2.xlsx", "RETAIL TRADE MONTHLY", "RETA.{c}.TOT.COF.BS.M", "ec_retail_conf",
         "Retail trade confidence indicator ((Q1-Q2+Q4)/3)", "balance, % points"),
        ("retail_total_sa_nace2.xlsx", "RETAIL TRADE MONTHLY", "RETA.{c}.TOT.6.BS.M", "ec_retail_sell_price_exp",
         "Retail trade (total): selling-price expectations next 3 months (Q6)", "balance, % points"),
        ("retail_subsectors_sa_m_nace2.xlsx", "FBT", "RETA.{c}.FBT.6.BS.M", "ec_food_retail_sell_price_exp",
         "Retail of food, beverages & tobacco (sub-sector FBT): selling-price expectations next 3 months (Q6)",
         "balance, % points"),
        ("industry_total_sa_nace2.xlsx", "INDUSTRY MONTHLY", "INDU.{c}.TOT.6.BS.M", "ec_ind_sell_price_exp",
         "Industry (total manufacturing): selling-price expectations for months ahead (Q6)", "balance, % points"),
        ("industry_subsectors_sa_m_nace2.xlsx", "10", "INDU.{c}.10.COF.BS.M", "ec_food_ind_conf",
         "Manufacture of food products (NACE C10): industrial confidence indicator ((Q2-Q4+Q5)/3)",
         "balance, % points"),
        ("industry_subsectors_sa_m_nace2.xlsx", "10", "INDU.{c}.10.6.BS.M", "ec_food_ind_sell_price_exp",
         "Manufacture of food products (NACE C10): selling-price expectations for months ahead (Q6)",
         "balance, % points"),
    ]
    cache = {}
    for fname, sheet, pat, suffix, desc, unit in specs:
        key = (fname, sheet)
        if key not in cache:
            cache[key] = ec_sheet(xdir, fname, sheet)
        df = cache[key]
        for cc, cname in EC_COUNTRIES.items():
            col = pat.format(c=cc)
            if col not in df.columns:
                FAILURES.append((f"{cc.lower()}_{suffix}", f"column {col} not in EC file {fname}"))
                continue
            notes = []
            if cc == "EE":
                notes.append("EC: Estonian surveys temporarily suspended from May 2026 (partner institute change)")
            if cc == "PL" and fname.startswith("consumer"):
                notes.append("EC: new Polish consumer-survey provider from 2018 (history replaced May 2026; "
                             "pre-2018 balances level-shifted by EC)")
            if cc == "EA":
                notes.append("Euro area aggregate as computed by EC (current composition)")
            if cc == "EU":
                notes.append("EU27 aggregate (current composition) as computed by EC")
            unit_ = unit
            add(f"{cc.lower()}_{suffix}", df[col], description=f"{cname}: {desc}", country=cc,
                frequency="M", unit=unit_, sa="SA", source=src,
                source_query=f"{EC_BASE}{folder}/{EC_ZIP_OF[fname]}.zip -> {fname} sheet '{sheet}' column {col}",
                lag=-2, notes="; ".join(notes + [lag_note]))


# ----------------------------------------------------------------------------- 2. OECD
OECD_URL = ("https://sdmx.oecd.org/public/rest/data/OECD.SDD.STES,DSD_STES@DF_CLI,/"
            "{areas}.M.{measures}.IX._Z.AA.IX._Z.H?startPeriod=1950-01&format=csvfile")
OECD_AREAS = {"NOR": ("no", "Norway"), "SWE": ("se", "Sweden"), "DNK": ("dk", "Denmark"),
              "FIN": ("fi", "Finland"), "CZE": ("cz", "Czechia"), "AUT": ("at", "Austria"),
              "POL": ("pl", "Poland"), "DEU": ("de", "Germany"), "IND": ("in", "India"),
              "OECD": ("oecd", "OECD total"), "EA20": ("ea", "Euro area (20)"),
              "G7": ("g7", "G7")}
OECD_MEAS = {"CCICP": ("oecd_cci", "OECD composite consumer confidence indicator (amplitude adjusted)"),
             "BCICP": ("oecd_bci", "OECD composite business confidence indicator (amplitude adjusted)"),
             "LI": ("oecd_cli", "OECD composite leading indicator (amplitude adjusted)")}


def fetch_oecd(dl: Path):
    url = OECD_URL.format(areas="+".join(OECD_AREAS), measures="+".join(OECD_MEAS))
    r = http_get(url, dl, "oecd_df_cli.csv", timeout=300)
    d = pd.read_csv(io.BytesIO(r.content))
    for (area, meas), g in d.groupby(["REF_AREA", "MEASURE"]):
        if area not in OECD_AREAS or meas not in OECD_MEAS:
            continue
        cc, cname = OECD_AREAS[area]
        suf, desc = OECD_MEAS[meas]
        s = pd.Series(g["OBS_VALUE"].values, index=pd.to_datetime(g["TIME_PERIOD"] + "-01"))
        notes = "OECD harmonised methodology; long-term average = 100. Revised each month."
        if meas == "BCICP" and area == "NOR":
            notes += " Norway BCI is derived from the quarterly SSB business tendency survey (lags 1 quarter)."
        if meas in ("CCICP", "BCICP") and area in ("SWE", "DNK", "FIN", "CZE", "AUT", "POL", "DEU", "EA20"):
            notes += " Based on the EC harmonised surveys (amplitude-adjusted transformation of EC balances)."
        add(f"{cc}_{suf}", s, description=f"{cname}: {desc}", country=cc.upper() if len(cc) == 2 else cc.upper(),
            frequency="M", unit="index, long-term average=100 (amplitude adjusted)", sa="SA",
            source="OECD SDMX API, dataflow OECD.SDD.STES:DSD_STES@DF_CLI(4.1)",
            source_query=f"key {area}.M.{meas}.IX._Z.AA.IX._Z.H", lag=12 if meas == "LI" else 15, notes=notes)
    for area in ("NOR",):
        if f"no_oecd_cci" not in SERIES:
            FAILURES.append(("no_oecd_cci", "OECD publishes no consumer confidence indicator for Norway "
                             "(use Finans Norge Forventningsbarometer no_fn_* instead)"))
    for area in ("NOR", "SWE", "DNK", "FIN", "CZE", "AUT", "POL", "OECD", "EA20"):
        cc = OECD_AREAS[area][0]
        if f"{cc}_oecd_cli" not in SERIES:
            FAILURES.append((f"{cc}_oecd_cli", "OECD CLI no longer published for this area "
                             "(since the 2023 CLI redesign the OECD computes CLIs only for G7/G20-type "
                             "economies and regional aggregates G7, G20, G4E, NAFTA, A5M); G7 CLI stored as global proxy"))


# ----------------------------------------------------------------------------- 3. UMich
def fetch_umich(dl: Path):
    src_q = "https://www.sca.isr.umich.edu/files/tbmics.csv (column ICS_ALL)"
    try:
        r = http_get("https://www.sca.isr.umich.edu/files/tbmics.csv", dl, "umich_tbmics.csv")
        d = pd.read_csv(io.BytesIO(r.content))
        d["date"] = pd.to_datetime(d["Month"].str.strip() + " " + d["YYYY"].astype(str), format="%B %Y")
        s = pd.Series(d["ICS_ALL"].values, index=d["date"])
        source = "University of Michigan Surveys of Consumers (identical to FRED UMCSENT)"
    except Exception as e:  # noqa: BLE001
        print("UMich direct failed, using FRED:", e)
        # FRED rejects some browser-like user agents through the proxy; use a plain UA.
        r = http_get("https://fred.stlouisfed.org/graph/fredgraph.csv?id=UMCSENT", dl, "fred_umcsent.csv",
                     headers={"User-Agent": "curl/8"})
        d = pd.read_csv(io.BytesIO(r.content))
        s = pd.Series(d.iloc[:, 1].values, index=pd.to_datetime(d.iloc[:, 0]))
        source = "FRED UMCSENT (University of Michigan)"
        src_q = "https://fred.stlouisfed.org/graph/fredgraph.csv?id=UMCSENT"
    s = pd.to_numeric(s, errors="coerce")
    s = s[s.index >= "1978-01-01"]
    add("us_umich_sentiment", s, description="United States: University of Michigan index of consumer sentiment",
        country="US", frequency="M", unit="index, 1966Q1=100", sa="NSA", source=source, source_query=src_q,
        lag=-2, notes="Stored from 1978-01 (monthly survey start; earlier data are quarterly/irregular). "
                      "Final reading published on the last Friday of the reference month; preliminary ~mid-month.")


# ----------------------------------------------------------------------------- 4. NIER
NIER = "https://statistik.konj.se/PxWeb/api/v1/en/KonjBar/"


def pxweb_query(url, sel: dict, headers=None) -> dict:
    q = {"query": [{"code": k, "selection": {"filter": "item", "values": v}} for k, v in sel.items()],
         "response": {"format": "json-stat2"}}
    r = requests.post(url, json=q, headers=headers or UA, timeout=120)
    r.raise_for_status()
    js = r.json()
    ids = js["id"]
    cats = []
    for i in ids:
        idx = js["dimension"][i]["category"]["index"]
        cats.append(sorted(idx, key=lambda c: idx[c]) if isinstance(idx, dict) else list(idx))
    out: dict = {}
    for n, combo in enumerate(itertools.product(*cats)):
        out.setdefault(combo[:-1], {})[combo[-1]] = js["value"][n]
    return out


def nier_period(p: str) -> pd.Timestamp:
    m = re.match(r"(\d{4})M(\d{2})", p)
    return pd.Timestamp(int(m.group(1)), int(m.group(2)), 1)


def nier_series(d: dict) -> pd.Series:
    s = pd.Series({nier_period(k): v for k, v in d.items() if v is not None}, dtype=float)
    return s.sort_index()


def fetch_nier(dl: Path):
    src = "Konjunkturinstitutet (NIER) Economic Tendency Survey, PxWeb API statistik.konj.se"
    lag = -5
    lagnote = "Published in the last week of the reference month."
    # (a) indicators (all seasonally adjusted, mean 100 / sd 10)
    tbl = "indikatorer/Indikatorm.px"
    d = pxweb_query(NIER + tbl, {"Indikator": ["KIFI", "bhus", "bhusmakro", "bhusmikro", "BDHAN", "B4711X", "BLIVS"]})
    (dl / "nier_indikatorm.json").write_text(json.dumps({"|".join(k): v for k, v in d.items()}))
    ind = {
        "KIFI": ("se_nier_eti", "Sweden: NIER Economic Tendency Indicator (Barometerindikatorn)"),
        "bhus": ("se_nier_cons_conf", "Sweden: NIER consumer confidence indicator (CCI)"),
        "bhusmakro": ("se_nier_cons_macro", "Sweden: NIER consumer macro index (view of Swedish economy)"),
        "bhusmikro": ("se_nier_cons_micro", "Sweden: NIER consumer micro index (own economy)"),
        "BDHAN": ("se_nier_retail_conf", "Sweden: NIER retail trade confidence indicator (NACE 45+47)"),
        "B4711X": ("se_nier_grocery_retail_conf",
                   "Sweden: NIER confidence indicator, retail sale of non-durable goods / grocery retail "
                   "(dagligvaruhandel, NACE 47.11+47.2)"),
        "BLIVS": ("se_nier_food_mfg_conf",
                  "Sweden: NIER confidence indicator, manufacture of food, beverages & tobacco (NACE 10-12)"),
    }
    for code, (sid, desc) in ind.items():
        add(sid, nier_series(d[(code,)]), description=desc, country="SE", frequency="M",
            unit="index, mean=100, sd=10", sa="SA", source=src,
            source_query=f"{NIER}{tbl} Indikator={code}", lag=lag,
            notes=lagnote + (" NACE 10-12 (no separate NACE 10 confidence in NIER database)." if code == "BLIVS" else ""))
    # (b) manufacturing: food industry selling-price expectations
    tbl = "ftgmanad/Barindm.px"
    meta = requests.get(NIER + tbl, headers=UA, timeout=60).json()
    bcode = meta["variables"][0]["code"]
    d = pxweb_query(NIER + tbl, {bcode: ["BLIVS"], "Fråga": ["202"], "Serie": ["S"]})
    add("se_nier_food_mfg_sell_price_exp", nier_series(d[("BLIVS", "202", "S")]),
        description="Sweden: NIER manufacture of food, beverages & tobacco (NACE 10-12): domestic selling-price "
                    "expectations next 3 months",
        country="SE", frequency="M", unit="net balance, % points", sa="SA", source=src,
        source_query=f"{NIER}{tbl} {bcode}=BLIVS, Fråga=202, Serie=S", lag=lag, notes=lagnote)
    # (c) trade: grocery retail price expectations
    tbl = "ftgmanad/Barhanm.px"
    meta = requests.get(NIER + tbl, headers=UA, timeout=60).json()
    bcode = meta["variables"][0]["code"]
    d = pxweb_query(NIER + tbl, {bcode: ["B4711X"], "Fråga": ["202", "250"], "Serie": ["S"]})
    add("se_nier_grocery_sell_price_exp", nier_series(d[("B4711X", "202", "S")]),
        description="Sweden: NIER grocery retail (NACE 47.11+47.2): selling-price expectations next 3 months",
        country="SE", frequency="M", unit="net balance, % points", sa="SA", source=src,
        source_query=f"{NIER}{tbl} {bcode}=B4711X, Fråga=202, Serie=S", lag=lag, notes=lagnote)
    add("se_nier_grocery_purch_price_exp", nier_series(d[("B4711X", "250", "S")]),
        description="Sweden: NIER grocery retail (NACE 47.11+47.2): purchase-price (goods) expectations next 3 months",
        country="SE", frequency="M", unit="net balance, % points", sa="SA", source=src,
        source_query=f"{NIER}{tbl} {bcode}=B4711X, Fråga=250, Serie=S", lag=lag,
        notes=lagnote + " Question introduced Oct 2015.")
    # (d) households' inflation expectations
    tbl = "hushall/hushallinfrant.px"
    d = pxweb_query(NIER + tbl, {"Fråga": ["Q061", "Q661"], "Serie": ["S", "T"], "Grupp": ["100"]})
    add("se_nier_hh_infl_exp_12m_mean", nier_series(d[("Q061", "T", "100")]),
        description="Sweden: households' expected inflation 12 months ahead (EU-harmonised quantitative question, "
                    "mean excl. extreme values)",
        country="SE", frequency="M", unit="percent", sa="NSA", source=src,
        source_query=f"{NIER}{tbl} Fråga=Q061, Serie=T, Grupp=100", lag=lag, notes=lagnote)
    add("se_nier_hh_infl_exp_12m_median", nier_series(d[("Q661", "S", "100")]),
        description="Sweden: households' expected inflation 12 months ahead (NIER question, median)",
        country="SE", frequency="M", unit="percent", sa="NSA", source=src,
        source_query=f"{NIER}{tbl} Fråga=Q661, Serie=S, Grupp=100", lag=lag,
        notes=lagnote + " Available from Apr 2015 only; use the _mean series for longer history.")


# ----------------------------------------------------------------------------- 5. SSB
def fetch_ssb(dl: Path):
    base = "https://data.ssb.no/api/v0/en/table/"
    src = "Statistics Norway (SSB), business tendency survey for manufacturing, mining & quarrying (Konjunkturbarometeret)"
    lag = 22
    lagnote = ("Quarterly; published ~3 weeks after quarter end (e.g. 2026Q2 on 2026-07-22). "
               "SSB publishes no food/beverage sub-index; 'consumer goods' (P126) is the closest grouping.")
    d = pxweb_query(base + "08267", {"PKoder": ["P105", "P126"], "Justering": ["S"],
                                     "ContentsCode": ["SammensattKonj"]})
    for code, sid, label in (("P105", "no_ssb_bts_mfg_conf", "manufacturing"),
                             ("P126", "no_ssb_bts_consgoods_conf", "consumer goods industries")):
        s = pd.Series({quarter_label_to_date(k): v for k, v in d[(code, "S", "SammensattKonj")].items()
                       if v is not None})
        add(sid, s, description=f"Norway: SSB business tendency survey, industrial confidence indicator, {label}",
            country="NO", frequency="Q", unit="net balance, % points (avg of orders, stocks (inverted), production expectations)",
            sa="SA", source=src, source_query=f"SSB table 08267: PKoder={code}, Justering=S, ContentsCode=SammensattKonj",
            lag=lag, notes=lagnote)
    d = pxweb_query(base + "08264", {"PKoder": ["P126"], "ReferansePeriode": ["01", "02"], "SerieType": ["02"],
                                     "Justering": ["S"], "ContentsCode": ["PriserHjemme", "Innsatspriser", "Lonnsomhet"]})
    for (ref, cont), sid, desc, extra in (
            (("02", "PriserHjemme"), "no_ssb_bts_consgoods_home_price_exp",
             "expected change in prices on home market next quarter, consumer goods industries", ""),
            (("02", "Innsatspriser"), "no_ssb_bts_consgoods_input_price_exp",
             "expected change in input (cost) prices next quarter, consumer goods industries",
             " Question introduced 2011Q4."),
            (("01", "Lonnsomhet"), "no_ssb_bts_consgoods_profitability",
             "change in profitability vs previous quarter, consumer goods industries",
             " Question introduced 2011Q4.")):
        raw = d[("P126", ref, "02", "S", cont)]
        s = pd.Series({quarter_label_to_date(k): v for k, v in raw.items() if v is not None})
        add(sid, s, description=f"Norway: SSB business tendency survey, {desc}", country="NO", frequency="Q",
            unit="net balance, % points", sa="SA", source=src,
            source_query=f"SSB table 08264: PKoder=P126, ReferansePeriode={ref}, SerieType=02 (balance), "
                         f"Justering=S, ContentsCode={cont}", lag=lag, notes=lagnote + extra)


# ----------------------------------------------------------------------------- 6. Norges Bank Regional Network
def nb_latest_rn_xlsx(dl: Path) -> tuple[str, str]:
    base = "https://www.norges-bank.no"
    html = http_get(base + "/en/news-events/publications/Regional-network-reports/").text
    links = set(re.findall(r'href="(/en/news-events/publications/Regional-network-reports/(\d{4})/(\d)-\d{4}/)"', html))
    if not links:
        raise RuntimeError("no regional network report links found")
    link = max(links, key=lambda t: (int(t[1]), int(t[2])))[0]
    page = http_get(base + link).text
    m = re.search(r'href="([^"]+\.xlsx[^"]*)"', page)
    if not m:
        raise RuntimeError(f"no xlsx on {link}")
    return base + m.group(1).replace("&amp;", "&"), link


def read_block_sheet(path: Path, sheet: str, header_row_of_date: int | None = None):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = [list(r) for r in wb[sheet].iter_rows(values_only=True)]
    return rows


def fetch_nb_rn(dl: Path):
    url, link = nb_latest_rn_xlsx(dl)
    print("RN xlsx:", url)
    http_get(url, dl, "nb_regional_network.xlsx")
    p = dl / "nb_regional_network.xlsx"
    src = "Norges Bank Regional Network (quarterly survey of ~1500 enterprises)"
    lag = -15
    lagnote = ("Survey round n refers to quarter n; report published ~2-3 weeks BEFORE the end of the reference "
               "quarter. Since Survey 2023/1 questions are quantitative (percent change); Norges Bank has "
               "back-calculated history to 2005 on the new basis. Dates = first day of the survey quarter.")

    def parse(sheet, date_hdr_row_text="Date"):
        rows = read_block_sheet(p, sheet)
        hi = next(i for i, r in enumerate(rows) if r and r[0] == date_hdr_row_text)
        hdr = rows[hi]
        # block label row (e.g. 'Current quarter' / 'Next quarter') = nearest non-empty row above hdr
        blk = [None] * len(hdr)
        for i in range(hi - 1, -1, -1):
            if any(isinstance(x, str) for x in rows[i][1:]):
                cur = None
                for j in range(len(hdr)):
                    v = rows[i][j] if j < len(rows[i]) else None
                    if j > 0 and isinstance(v, str):
                        cur = v.strip()
                    blk[j] = cur
                break
        data = []
        for r in rows[hi + 1:]:
            if r[0] is None or not hasattr(r[0], "year"):
                continue
            data.append(r)
        return hdr, blk, data

    def col_series(hdr, blk, data, colname, block=None):
        for j, h in enumerate(hdr):
            if isinstance(h, str) and h.replace("-", "").replace(" ", "").lower() == colname.replace("-", "").replace(" ", "").lower() \
                    and (block is None or (blk[j] or "").lower().startswith(block.lower())):
                s = pd.Series({pd.Timestamp(r[0]).to_period("Q").to_timestamp(): r[j] for r in data})
                return s
        raise KeyError(f"{colname} / {block} not found")

    hdr, blk, data = parse("Output")
    for col, tag in (("Aggregated", "total"), ("Retail trade", "retail"), ("Household services", "hhserv")):
        for block, btag, bdesc in (("Current quarter", "cur", "current quarter (past 3 months)"),
                                   ("Next quarter", "next", "expected next quarter")):
            s = col_series(hdr, blk, data, col, block)
            add(f"no_nbrn_output_{btag}_{tag}", s,
                description=f"Norway: Norges Bank Regional Network, output growth {bdesc}, {col.lower()}",
                country="NO", frequency="Q", unit="percent, quarter-on-quarter (contacts' reported/expected)",
                sa="NSA", source=src,
                source_query=f"{url} sheet 'Output', block '{block}', column '{col}'", lag=lag,
                notes=lagnote + (" Household services sub-series = 'Household services' column." if tag == "hhserv" else ""))
    hdr, blk, data = parse("Profitability")
    for col, tag in (("Aggregated", "total"), ("Retail trade", "retail")):
        s = col_series(hdr, blk, data, col)
        add(f"no_nbrn_profit_{tag}", s,
            description=f"Norway: Norges Bank Regional Network, profitability: change in operating margin in current "
                        f"quarter vs same quarter a year earlier, {col.lower()}",
            country="NO", frequency="Q", unit="percent (as published by Norges Bank; net/diffusion-type measure)",
            sa="NSA", source=src, source_query=f"{url} sheet 'Profitability', column '{col}'", lag=lag, notes=lagnote)


# ----------------------------------------------------------------------------- 7. Norges Bank Expectations Survey
def fetch_nb_es(dl: Path):
    base = "https://www.norges-bank.no"
    html = http_get(base + "/en/topics/Monetary-policy/expectations-survey/").text
    m = re.search(r'href="([^"]+expectations_survey[^"]*\.xlsx[^"]*)"', html, re.I) or \
        re.search(r'href="([^"]+\.xlsx[^"]*)"', html)
    if not m:
        raise RuntimeError("no expectations survey xlsx link found")
    url = base + m.group(1).replace("&amp;", "&") if m.group(1).startswith("/") else m.group(1)
    print("ES xlsx:", url)
    http_get(url, dl, "nb_expectations_survey.xlsx")
    p = dl / "nb_expectations_survey.xlsx"
    src = "Norges Bank Expectations Survey (conducted by Ipsos; quarterly)"
    lag = -40
    lagnote = ("Survey conducted early in the middle month of the quarter and published ~6 weeks before quarter end "
               "(e.g. 2026Q3 file dated 2026-08-20). Dates = first day of survey quarter.")

    def header_index(rows):
        n = max(len(r) for r in rows[:6])
        def ffill(row):
            out, cur = [], None
            for j in range(n):
                v = row[j] if j < len(row) else None
                if isinstance(v, str) and v.strip():
                    cur = re.sub(r"\s+", " ", v).strip()
                out.append(cur)
            return out
        top, grp, sub = ffill(rows[0]), ffill(rows[1]), ffill(rows[3])
        starts = [j for j in range(n) if j < len(rows[0]) and rows[0][j]]
        bstart = []
        for j in range(n):
            st = max([s for s in starts if s <= j], default=0)
            bstart.append(st)
        # question number per block: first 'Question NN' text in [start, next start)
        qnum = {}
        for k, st in enumerate(starts):
            en = starts[k + 1] if k + 1 < len(starts) else n
            q = None
            for j in range(st, en):
                v = rows[2][j] if j < len(rows[2]) else None
                if isinstance(v, str) and re.search(r"Questi?on\s*(\d+)", v):
                    q = int(re.search(r"Questi?on\s*(\d+)", v).group(1)); break
            qnum[st] = q
        cols = []
        for j in range(n):
            stat = rows[4][j] if j < len(rows[4]) and rows[4][j] else None
            if stat is None and j < len(rows[3]) and isinstance(rows[3][j], str) and rows[3][j].strip() in (
                    "Average", "Diffusion index", "Weighted median", "Unweighted median"):
                stat = rows[3][j].strip()
            cols.append(dict(j=j, top=top[j], grp=grp[j], sub=sub[j], q=qnum.get(bstart[j]),
                             stat=re.sub(r"\s+", " ", stat).strip() if isinstance(stat, str) else None))
        return cols

    def pick(rows, cols, **cond):
        hits = []
        for c in cols:
            ok = True
            for k, v in cond.items():
                cv = c[k]
                if k == "q":
                    ok &= cv == v
                elif k == "top":
                    ok &= cv is not None and v.lower() in cv.lower()
                else:
                    ok &= cv is not None and cv.strip().lower() == v.lower()
            if ok:
                hits.append(c["j"])
        if len(hits) != 1:
            raise KeyError(f"ambiguous/missing column for {cond}: {hits}")
        j = hits[0]
        s = {}
        for r in rows[5:]:
            d = quarter_label_to_date(r[0]) if r and r[0] else None
            if d is not None and j < len(r) and isinstance(r[j], (int, float)):
                s[d] = r[j]
        return pd.Series(s), j

    import openpyxl
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    inf = [list(r) for r in wb["INFLATION"].iter_rows(values_only=True)]
    ic = header_index(inf)
    emp_sheet = next(s for s in wb.sheetnames if "PROFITABILITY" in s.upper())
    emp = [list(r) for r in wb[emp_sheet].iter_rows(values_only=True)]
    ec = header_index(emp)
    specs = [
        ("no_nbes_hh_infl_exp_12m", inf, ic, dict(top="NEXT 12 MONTHS", grp="HOUSEHOLDS", q=32, stat="Average"),
         "Norway: households' expected CPI inflation 12 months ahead (mean)", "percent",
         "Q32 'About how much higher/lower' (mean). Households series from 2002Q3."),
        ("no_nbes_hh_infl_exp_2_3y", inf, ic, dict(top="TWO YEARS", grp="HOUSEHOLDS", q=33, stat="Average"),
         "Norway: households' expected CPI inflation in 2-3 years (mean)", "percent", "Q33."),
        ("no_nbes_bl_infl_exp_12m", inf, ic, dict(top="NEXT 12 MONTHS", grp="BUSINESS LEADERS", q=19,
                                                  sub="Business leaders", stat="Average"),
         "Norway: business leaders' expected CPI inflation 12 months ahead (mean)", "percent", "Q19."),
        ("no_nbes_bl_infl_exp_2y", inf, ic, dict(top="TWO YEARS", grp="BUSINESS LEADERS", q=20,
                                                 sub="Business leaders", stat="Average"),
         "Norway: business leaders' expected CPI inflation in 2 years (mean)", "percent", "Q20."),
        ("no_nbes_econ_infl_exp_2y", inf, ic, dict(top="TWO YEARS", grp="ECONOMISTS", sub="Economists", stat="Average"),
         "Norway: economists' (financial industry + academia) expected CPI inflation in 2 years (mean)", "percent", "Q2."),
        ("no_nbes_bl_purch_price_di", inf, ic, dict(top="NEXT 12 MONTHS", grp="BUSINESS LEADERS", q=22,
                                                    sub="Business leaders", stat="Diffusion index"),
         "Norway: business leaders' expected purchase-price growth next 12m vs past 12m (diffusion index)",
         "diffusion index (50 = same pace)", "Q22: increase more / same pace / less than past 12 months."),
        ("no_nbes_bl_sell_price_di", inf, ic, dict(top="NEXT 12 MONTHS", grp="BUSINESS LEADERS", q=23,
                                                   sub="Business leaders", stat="Diffusion index"),
         "Norway: business leaders' expected selling-price growth next 12m vs past 12m (diffusion index)",
         "diffusion index (50 = same pace)", "Q23: increase more / same pace / less than past 12 months."),
        ("no_nbes_bl_profit_next12m", emp, ec, dict(top="PROFITABILITY", q=27, sub="Business leaders",
                                                    stat="Profitability index"),
         "Norway: business leaders' expected change in profitability (EBITDA margin) next 12 months (profitability index)",
         "net balance (improve minus weaken), % points", "Q27."),
        ("no_nbes_bl_profit_past12m", emp, ec, dict(top="PROFITABILITY", q=26, sub="Business leaders",
                                                    stat="Profitability index"),
         "Norway: business leaders' reported change in profitability (EBITDA margin) past 12 months (profitability index)",
         "net balance (improved minus weakened), % points", "Q26."),
    ]
    for sid, rows, cols, cond, desc, unit, note in specs:
        s, j = pick(rows, cols, **cond)
        add(sid, s, description=desc, country="NO", frequency="Q", unit=unit, sa="NSA", source=src,
            source_query=f"{url} sheet '{'INFLATION' if rows is inf else emp_sheet}' column index {j} ({cond})",
            lag=lag, notes=note + " " + lagnote)


# ----------------------------------------------------------------------------- 8. Finans Norge
FN_PAGE = "https://www.finansnorge.no/tema/statistikk-og-analyse/forventningsbarometeret/siste-kvartalstall/"


def dw_latest_csv(chart_id: str, dl: Path) -> tuple[pd.DataFrame, str]:
    r = http_get(f"https://datawrapper.dwcdn.net/{chart_id}/")
    m = re.search(r"datawrapper\.dwcdn\.net/%s/(\d+)/" % chart_id, r.text)
    ver = m.group(1) if m else None
    url = f"https://datawrapper.dwcdn.net/{chart_id}/{ver}/dataset.csv" if ver else \
        f"https://datawrapper.dwcdn.net/{chart_id}/dataset.csv"
    r = http_get(url, dl, f"dw_{chart_id}_{ver}.csv")
    txt = r.content.decode("utf-8-sig")
    sep = "\t" if txt.split("\n")[0].count("\t") > txt.split("\n")[0].count(",") else ","
    df = pd.read_csv(io.StringIO(txt), sep=sep)
    return df, url


def fetch_finansnorge(dl: Path):
    html = http_get(FN_PAGE, dl, "finansnorge_siste.html").text
    ids = list(dict.fromkeys(re.findall(r"datawrapper\.dwcdn\.net/([A-Za-z0-9]{5})/", html)))
    for fb in ("dPMwa", "4EDKP"):
        if fb not in ids:
            ids.append(fb)
    main = sub = None
    for cid in ids:
        try:
            df, url = dw_latest_csv(cid, dl)
        except Exception as e:  # noqa: BLE001
            print("datawrapper", cid, e)
            continue
        cols = [c.lower() for c in df.columns]
        if any("justert" in c and "ujustert" not in c for c in cols) and main is None:
            main = (df, url, cid)
        elif any(c.startswith("egen") for c in cols) and any(c.startswith("landet") for c in cols) and sub is None:
            sub = (df, url, cid)
    if main is None or sub is None:
        raise RuntimeError("Finans Norge datawrapper charts not identified")
    src = "Finans Norge / Verian (Kantar) Forventningsbarometeret (household expectations survey), finansnorge.no"
    lag = -45
    lagnote = ("Quarterly household survey (since 1992Q3). Interviews in the first weeks "
               "of the middle month of the quarter; results published ~mid-quarter (2026Q3 on 2026-08-12). "
               "Data scraped from the Datawrapper chart behind finansnorge.no 'Siste kvartalstall'.")

    def qser(df, col):
        return pd.Series({quarter_label_to_date(a): b for a, b in zip(df.iloc[:, 0], df[col])
                          if quarter_label_to_date(a) is not None})

    df, url, cid = main
    cmap = {c.lower(): c for c in df.columns}
    adj = next(cmap[c] for c in cmap if "justert" in c and "ujustert" not in c)
    unadj = next(cmap[c] for c in cmap if "ujustert" in c)
    add("no_fn_cci_sa", qser(df, adj),
        description="Norway: Forventningsbarometeret main indicator (household consumer confidence), seasonally adjusted",
        country="NO", frequency="Q", unit="net balance, % points (avg of 5 sub-indicators)", sa="SA", source=src,
        source_query=f"{url} column '{adj}'", lag=lag,
        notes=lagnote + " SA series is re-estimated each quarter (latest vintage stored).")
    add("no_fn_cci_nsa", qser(df, unadj),
        description="Norway: Forventningsbarometeret main indicator (household consumer confidence), unadjusted",
        country="NO", frequency="Q", unit="net balance, % points (avg of 5 sub-indicators)", sa="NSA", source=src,
        source_query=f"{url} column '{unadj}'", lag=lag, notes=lagnote)
    df, url, cid = sub
    want = [("egen", "neste", "no_fn_own_econ_next_yr", "own household economy next year"),
            ("egen", "siste", "no_fn_own_econ_past_yr", "own household economy past year"),
            ("landet", "neste", "no_fn_country_econ_next_yr", "Norway's economy next year"),
            ("landet", "siste", "no_fn_country_econ_past_yr", "Norway's economy past year"),
            ("større", "anskaff", "no_fn_big_purchases", "good time for larger purchases")]
    for a, b, sid, desc in want:
        col = next(c for c in df.columns if c.lower().startswith(a) and b in c.lower())
        add(sid, qser(df, col), description=f"Norway: Forventningsbarometeret sub-indicator: {desc}",
            country="NO", frequency="Q", unit="net balance, % points", sa="NSA", source=src,
            source_query=f"{url} column '{col}'", lag=lag, notes=lagnote)


# ----------------------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--download-dir", default=None)
    a = ap.parse_args()
    dl = Path(a.download_dir) if a.download_dir else Path(tempfile.mkdtemp(prefix="surveys_dl_"))
    dl.mkdir(parents=True, exist_ok=True)
    print("download dir:", dl)
    run_source("EC BCS", fetch_ec, dl)
    run_source("OECD", fetch_oecd, dl)
    run_source("UMich", fetch_umich, dl)
    run_source("NIER", fetch_nier, dl)
    run_source("SSB", fetch_ssb, dl)
    run_source("Norges Bank Regional Network", fetch_nb_rn, dl)
    run_source("Norges Bank Expectations Survey", fetch_nb_es, dl)
    run_source("Finans Norge", fetch_finansnorge, dl)
    write_all()
    print(f"\nwrote {len(SERIES)} series to {OUT}")
    if FAILURES:
        print("\nFAILURES / not available:")
        for k, v in FAILURES:
            print(f"  {k}: {v}")
    (dl / "_fetch_failures.json").write_text(json.dumps(FAILURES, indent=1))


if __name__ == "__main__":
    main()
